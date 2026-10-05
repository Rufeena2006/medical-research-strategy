"""Build the Week 1 research-strategy Word document (.docx).

Usage:
    python -m src.generate_report --author "Your Name"

If outputs/ contains results from `python -m src.run_search`, live search
counts, the publication-trend chart and the screening summary are inserted
automatically; otherwise a clearly marked placeholder is shown.
"""
import argparse
import csv
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import config

NAVY = RGBColor(0x1F, 0x3A, 0x5F)


# ---------- helpers ----------------------------------------------------------
def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para(doc, text, bold=False, italic=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(2)


def numbered(doc, items):
    for n, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{n}.\t")
        if isinstance(it, tuple):
            r = p.add_run(it[0])
            r.bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(0.3))


def table(doc, headers, rows, widths, font=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "1F3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(font)
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa")
    tblW.set(qn("w:w"), str(int(sum(widths) * 1440)))
    for i, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] * 1440)))
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def h(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for r in heading.runs:
        r.font.color.rgb = NAVY
    return heading


# ---------- document ---------------------------------------------------------
def build(author: str, out_dir: Path, outfile: Path):
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    # Title block
    para(doc, "Week 1 Task: Research Planning and Strategy Development", bold=True, size=20,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4).runs[0].font.color.rgb = NAVY
    para(doc, "Medical Writing Internship: Research Strategy Document", size=13,
         align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    para(doc, config.TOPIC_TITLE, bold=True, size=13, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    table(doc, ["Item", "Detail"], [
        ["Prepared by", author],
        ["Date", date.today().strftime("%d %B %Y")],
        ["Code repository", "glp1-obesity-literature-strategy (search code, search log, this report)"],
        ["Data sources", "Publicly available only: PubMed, Europe PMC, Google Scholar, ClinicalTrials.gov, "
                         "Cochrane Library abstracts, medRxiv, regulatory (FDA/EMA) public documents"],
    ], [1.5, 5.0], font=10)

    # 1 Topic and rationale
    h(doc, "1. Research Topic and Rationale")
    para(doc, "Topic: Incretin-based therapies (semaglutide, tirzepatide, liraglutide and emerging agents) for "
              "chronic weight management in adults with obesity or overweight: how much weight they remove, what "
              "they do for heart and kidney outcomes, how safe they are, and what happens when treatment stops.")
    para(doc, "Why this topic?", bold=True, space_after=2)
    bullets(doc, [
        ("Clinical importance: ", "obesity is a chronic, relapsing disease linked to type 2 diabetes, cardiovascular "
                                  "disease, kidney disease and several cancers; effective pharmacotherapy has been "
                                  "scarce until recently."),
        ("Contemporary relevance: ", "semaglutide 2.4 mg and tirzepatide have been approved for chronic weight "
                                     "management in recent years, and new agents (e.g. retatrutide, orforglipron) are "
                                     "in late-stage trials, so the evidence base is moving quickly."),
        ("Adequate public literature: ", "pivotal RCTs, meta-analyses, regulatory reviews and growing real-world "
                                         "cohorts are available through free, open resources."),
        ("Medical-writing value: ", "the topic supports several later writing outputs (evidence summary, "
                                    "lay summary, safety overview, slide deck) because it combines efficacy data, "
                                    "safety data and patient-facing communication needs."),
        ("Feasibility: ", "a well-defined drug class and outcome set keeps the review scoped to a 30 to 35 hour effort."),
    ])

    # 2 State of research
    h(doc, "2. Summary of the Current State of Research")
    para(doc, "Preliminary scoping of landmark publications (to be re-verified against PubMed during Step 2 of the "
              "workflow) shows the following picture.")
    table(doc, ["Domain", "Key evidence (illustrative landmark sources)", "Take-away"], [
        ["Efficacy: semaglutide",
         "STEP 1 RCT (Wilding et al., N Engl J Med 2021): semaglutide 2.4 mg weekly vs placebo for 68 weeks "
         "in adults with overweight/obesity without diabetes.",
         "Mean weight loss of roughly 15% vs about 2% with placebo."],
        ["Efficacy: tirzepatide",
         "SURMOUNT-1 RCT (Jastreboff et al., N Engl J Med 2022): once-weekly tirzepatide for 72 weeks.",
         "Dose-dependent loss of about 15% to 21% vs about 3% with placebo."],
        ["Head-to-head",
         "SURMOUNT-5 (Aronne et al., N Engl J Med 2025): tirzepatide vs semaglutide.",
         "Greater weight reduction with tirzepatide; check final figures in source."],
        ["Cardiovascular",
         "SELECT RCT (Lincoff et al., N Engl J Med 2023): semaglutide in adults with established CVD and "
         "overweight/obesity without diabetes.",
         "Lower rate of major adverse cardiovascular events (about 20% relative reduction)."],
        ["Durability / withdrawal",
         "STEP 1 extension (Wilding et al., Diabetes Obes Metab 2022); SURMOUNT-4 (Aronne et al., JAMA 2024).",
         "Substantial weight regain after stopping; continued treatment maintains loss."],
        ["Safety",
         "Pooled trial safety data and regulatory labels.",
         "Gastrointestinal adverse events are most common; rarer signals (gallbladder disease, "
         "pancreatitis, others) need careful synthesis."],
    ], [1.2, 3.3, 2.0])
    para(doc, "Overall: efficacy and cardiovascular evidence in selected populations is strong; evidence is thinner "
              "on long-term durability, post-discontinuation outcomes, real-world persistence and equity of access.",
         italic=True)

    # 3 Gaps
    h(doc, "3. Gaps Identified in the Literature")
    bullets(doc, [
        "Long-term (more than 2 years) efficacy and safety outside highly selected trial populations.",
        "Weight regain and cardiometabolic rebound after discontinuation, and how common stopping is in routine care.",
        "Real-world persistence, adherence and dose patterns compared with trial results.",
        "Head-to-head comparisons between agents and evidence on switching strategies.",
        "Representation of older adults, people with severe comorbidity and under-represented ethnic groups.",
        "Consistency of reporting for rare safety signals across trials, cohorts and regulatory sources.",
    ])

    # 4 Research questions
    h(doc, "4. Research Questions and Hypotheses")
    table(doc, ["#", "Research question", "Working hypothesis"], [
        ["RQ1", "In adults with obesity, how large and how durable is weight loss with semaglutide or tirzepatide "
                "compared with placebo or each other?",
         "Tirzepatide produces greater mean weight loss than semaglutide; loss plateaus after about 60 to 72 weeks."],
        ["RQ2", "What evidence exists for cardiovascular and kidney benefit beyond weight loss?",
         "Benefit is demonstrated mainly in populations with established cardiovascular disease; evidence in "
         "primary prevention is limited."],
        ["RQ3 (gap)", "What happens to weight and cardiometabolic markers after treatment is stopped, and how "
                      "often do patients stop in real-world settings?",
         "Most lost weight is regained within one to two years of stopping, and real-world discontinuation exceeds "
         "trial rates."],
        ["RQ4", "How consistently are gastrointestinal and rare safety events reported across study designs?",
         "Common events are consistent; rare events are inconsistently defined and reported."],
    ], [0.8, 3.1, 2.6])

    # 5 Search strategy
    h(doc, "5. Search Strategy")
    h(doc, "5.1 Databases and sources", 2)
    table(doc, ["Source", "Purpose", "Access"], [
        ["PubMed / MEDLINE", "Primary biomedical database; Boolean + MeSH searching; searched programmatically", "Free"],
        ["Europe PMC", "Cross-check, preprints, full-text open access", "Free"],
        ["Google Scholar", "Citation chasing, grey literature, backward/forward snowballing", "Free"],
        ["ClinicalTrials.gov / WHO ICTRP", "Registered and ongoing trials; unpublished results", "Free"],
        ["Cochrane Library (abstracts)", "Existing systematic reviews", "Free abstracts"],
        ["medRxiv", "Recent preprints (flagged as non-peer-reviewed)", "Free"],
        ["FDA / EMA public documents", "Labels, review documents, safety communications", "Free"],
    ], [1.9, 3.5, 1.1])

    h(doc, "5.2 Concept framework (PICO)", 2)
    table(doc, ["Element", "Definition"], [
        ["Population", "Adults (18 years or older) with obesity, or overweight with at least one weight-related "
                       "comorbidity; with or without type 2 diabetes (diabetes analysed as subgroup)"],
        ["Intervention", "Semaglutide, tirzepatide, liraglutide, and other incretin-based agents for weight management"],
        ["Comparator", "Placebo, lifestyle intervention, active comparator, or treatment discontinuation"],
        ["Outcomes", "Percent body-weight change; cardiovascular and kidney outcomes; adverse events; weight regain "
                     "after stopping; persistence and adherence"],
        ["Study designs", "RCTs, systematic reviews/meta-analyses, large observational cohorts and real-world studies"],
    ], [1.3, 5.2])

    h(doc, "5.3 Preliminary search terms", 2)
    table(doc, ["Concept", "Controlled vocabulary and free-text terms"], [
        ["Population", "\"Obesity\"[Mesh]; \"Overweight\"[Mesh]; obesity; obese; overweight; weight management; weight loss"],
        ["Intervention", "\"Glucagon-Like Peptide-1 Receptor\"[Mesh]; semaglutide; tirzepatide; liraglutide; "
                         "retatrutide; orforglipron; GLP-1; glucagon-like peptide-1; incretin"],
        ["Cardio-renal", "cardiovascular; major adverse cardiovascular events; MACE; heart failure; mortality; "
                         "chronic kidney disease"],
        ["Durability", "discontinu*; withdraw*; weight regain; weight maintenance; cessation; persistence; adherence"],
        ["Safety", "adverse events; safety; pancreatitis; gallbladder; cholelithiasis; gastrointestinal; "
                   "optic neuropathy; suicid*"],
        ["Real-world", "real-world; electronic health records; retrospective cohort; claims data; target trial emulation"],
        ["Design filter", "Randomized Controlled Trial[pt]; Meta-Analysis[pt]; Systematic Review[pt]; "
                          "Cohort Studies[Mesh]; Observational Study[pt]"],
    ], [1.3, 5.2])

    h(doc, "5.4 Search logic and reproducibility", 2)
    bullets(doc, [
        "Terms within a concept are combined with OR; concepts are combined with AND (Population AND Intervention is the core).",
        "Named queries Q1 to Q7 combine the core with each outcome block; the full strings are in Appendix A and in src/config.py.",
        "Limits: human studies, English language, publication date from 1 January 2019 to the search date.",
        "The search is executed by code (NCBI E-utilities), and each run is logged with the date, query and hit count.",
        "Backward and forward citation chasing of included landmark papers via Google Scholar and PubMed 'Similar articles'.",
        "Search is repeated before final writing to capture new publications.",
    ])

    # 6 Selection criteria
    h(doc, "6. Article Selection Criteria")
    table(doc, ["Inclusion", "Exclusion"], [
        ["Adults 18 years or older with obesity or overweight", "Animal, in vitro or purely mechanistic studies"],
        ["GLP-1 receptor agonist or dual/triple incretin agonist used for weight management",
         "Paediatric-only populations (planned as a separate review)"],
        ["Reports weight, cardiometabolic, safety, discontinuation or persistence outcomes",
         "Editorials, letters, comments, case reports, news items"],
        ["RCT, systematic review/meta-analysis, or cohort/real-world study with at least 100 participants "
         "(observational)", "No abstract or no accessible full text"],
        ["Published 2019 onwards, English language", "Non-English; duplicate or retracted publications"],
        ["Peer-reviewed (preprints used only as flagged supplements)", "Glycaemic-only endpoints with no weight outcome"],
    ], [3.25, 3.25])

    # 7 Workflow
    h(doc, "7. Screening, Extraction and Appraisal Workflow")
    numbered(doc, [
        ("Run searches and log results. ", "Execute src/run_search.py; save counts, strings and dates."),
        ("De-duplicate and rule-based pre-screen. ", "The script flags clear exclusions (wrong population, "
                                                      "publication type, language); a human confirms every decision."),
        ("Title/abstract screening. ", "Candidate list reviewed manually (Zotero and a spreadsheet or Rayyan free tier)."),
        ("Full-text review. ", "Open-access full texts read against the criteria in Section 6; reasons for exclusion recorded."),
        ("Data extraction. ", "Standard sheet: design, population, intervention, comparator, duration, "
                              "primary outcome, key numbers, safety events, limitations."),
        ("Quality appraisal. ", "Cochrane RoB 2 for RCTs, AMSTAR 2 for reviews, Newcastle-Ottawa Scale for cohorts."),
        ("Synthesis. ", "Narrative synthesis with evidence tables; meta-analytic figures are only quoted from "
                        "published pooled analyses."),
        ("Reporting. ", "PRISMA-style flow diagram built from prisma_summary.json."),
    ])

    # 8 Writing plan
    h(doc, "8. Medical Writing Plan for Later Tasks")
    bullets(doc, [
        "Structured evidence summary organised by RQ1 to RQ4, with an evidence table and a limitations section.",
        "Plain-language summary for patients and a short clinician-focused key messages page.",
        "Consistent citation style (Vancouver) managed in Zotero; every numeric claim traceable to a source.",
        "Balanced wording: efficacy and safety presented together; uncertainty and evidence quality stated explicitly.",
    ])

    # 9 Feasibility
    h(doc, "9. Feasibility, Timeline and Risk Management")
    table(doc, ["Activity", "Estimated hours"], [
        ["Topic selection and scoping", "3"],
        ["Preliminary literature review of landmark papers", "6"],
        ["Search strategy design, testing and refinement", "6"],
        ["Screening of candidate records", "6"],
        ["Evidence extraction and synthesis of current state", "5"],
        ["Code, testing and repository set-up", "4"],
        ["Drafting and formatting this document", "5"],
        ["Total", "35"],
    ], [4.5, 2.0])
    table(doc, ["Risk", "Mitigation"], [
        ["Too many records", "Use design filters, date limits and the Q7 retrieval set; prioritise RCTs and systematic reviews."],
        ["Paywalled full text", "Use PMC, Europe PMC open access, author manuscripts and abstracts; note limitation."],
        ["Rapidly changing field", "Re-run the logged searches before final writing."],
        ["Reviewer bias", "Predefined criteria, recorded exclusion reasons, appraisal tools."],
        ["API limits or outages", "Rate-limited client with retries; optional NCBI API key."],
    ], [2.0, 4.5])

    # 10 Preliminary yield
    h(doc, "10. Preliminary Search Yield")
    counts_file = out_dir / "search_counts.csv"
    if counts_file.exists():
        with counts_file.open(encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        if rows:
            para(doc, f"Search run on {rows[0]['run_at']} (PubMed, {rows[0]['date_limit']}).")
            table(doc, ["Query", "Hits"], [[r["query"], r["hits"]] for r in rows], [4.5, 2.0])
        png = out_dir / "yearly_counts.png"
        if png.exists():
            doc.add_picture(str(png), width=Inches(5.6))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            para(doc, "Figure 1. Publication trend for the core query.", italic=True, size=9,
                 align=WD_ALIGN_PARAGRAPH.CENTER)
        pj = out_dir / "prisma_summary.json"
        if pj.exists():
            p = json.loads(pj.read_text(encoding="utf-8"))
            table(doc, ["Screening step", "Records"], [
                ["Matching the retrieval query (total)", p["records_matching_query_total"]],
                ["Downloaded for screening", p["records_downloaded"]],
                ["Excluded by rule-based pre-screen", p["excluded_by_rules"]],
                ["Candidates for manual review", p["include_candidates_for_manual_review"]],
            ], [4.5, 2.0])
    else:
        para(doc, "[Placeholder: run `python -m src.run_search` and then `python -m src.generate_report` "
                  "to insert live PubMed counts, the publication-trend chart and the screening summary here.]",
             italic=True)

    # 11 Reproducibility
    h(doc, "11. Reproducibility and Code Repository")
    bullets(doc, [
        "src/config.py: all search blocks and screening vocabulary.",
        "src/pubmed_client.py: PubMed E-utilities client; src/screening.py: rule-based pre-screening.",
        "src/run_search.py: runs the strategy and writes the search log, charts and screened records.",
        "src/generate_report.py: regenerates this document; tests/ contains offline unit tests.",
    ])

    # 12 References
    h(doc, "12. Key References (to be verified in PubMed before submission)")
    numbered(doc, [
        "Wilding JPH, et al. Once-weekly semaglutide in adults with overweight or obesity (STEP 1). N Engl J Med. 2021.",
        "Jastreboff AM, et al. Tirzepatide once weekly for the treatment of obesity (SURMOUNT-1). N Engl J Med. 2022.",
        "Lincoff AM, et al. Semaglutide and cardiovascular outcomes in obesity without diabetes (SELECT). N Engl J Med. 2023.",
        "Wilding JPH, et al. Weight regain and cardiometabolic effects after withdrawal of semaglutide: STEP 1 trial extension. Diabetes Obes Metab. 2022.",
        "Aronne LJ, et al. Continued treatment with tirzepatide for maintenance of weight reduction (SURMOUNT-4). JAMA. 2024.",
        "Aronne LJ, et al. Tirzepatide as compared with semaglutide for the treatment of obesity (SURMOUNT-5). N Engl J Med. 2025.",
        "Page MJ, et al. The PRISMA 2020 statement. BMJ. 2021.",
        "Sterne JAC, et al. RoB 2: a revised tool for assessing risk of bias in randomised trials. BMJ. 2019.",
    ])

    # Appendix
    doc.add_page_break()
    h(doc, "Appendix A. Full Search Strings (PubMed)")
    for name, q in config.QUERIES.items():
        para(doc, name, bold=True, space_after=1)
        p = para(doc, q, size=8.5, space_after=8)
        for r in p.runs:
            r.font.name = "Consolas"

    outfile.parent.mkdir(parents=True, exist_ok=True)
    doc.save(outfile)
    return outfile


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--author", default="[Your Name]")
    ap.add_argument("--outputs", default="outputs")
    ap.add_argument("--file", default="report/Week1_Research_Strategy_GLP1_Obesity.docx")
    args = ap.parse_args()
    path = build(args.author, Path(args.outputs), Path(args.file))
    print(f"Report written to {path}")


if __name__ == "__main__":
    main()
